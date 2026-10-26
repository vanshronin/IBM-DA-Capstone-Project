import pptx

prs = pptx.Presentation(r'TEMP/Data Analyst Capstone Template 2026.pptx')
print(f"Width: {prs.slide_width.inches:.2f} in, Height: {prs.slide_height.inches:.2f} in")
for idx in range(len(prs.slides)):
    slide = prs.slides[idx]
    title_text = slide.shapes.title.text if slide.shapes.title else "No Title"
    print(f"\n--- Slide {idx+1}: {title_text} ---")
    for s_idx, shape in enumerate(slide.shapes):
        print(f"  Shape {s_idx}: name='{shape.name}', type={shape.shape_type}, pos=({shape.left.inches:.2f}, {shape.top.inches:.2f}, {shape.width.inches:.2f}, {shape.height.inches:.2f})")
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    print(f"     p: {p.text.strip()[:90]}")
